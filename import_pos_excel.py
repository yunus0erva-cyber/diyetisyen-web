"""
Florya MEV Koleji - POS Cihazları Günlük Takip Veri Aktarım Aracı
Masaüstündeki 'Kantin POS Günlük Kasa Takibi (Seri No Bazlı).xlsx' dosyasındaki verileri aktarır.
"""

import openpyxl
from pathlib import Path
from canteen_manager import add_or_update_pos_record
from database import trigger_auto_backup

EXCEL_PATH = Path(r"C:\Users\doruk\Desktop\Kantin POS Günlük Kasa Takibi (Seri No Bazlı).xlsx")

def import_pos_data():
    if not EXCEL_PATH.exists():
        print(f"[HATA] Dosya bulunamadı: {EXCEL_PATH}")
        return False

    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    ws = wb.active

    count = 0
    # Satır 2'den 31'e kadar (1 Eylül - 30 Eylül)
    for r in range(2, 32):
        d_val = ws.cell(r, 1).value
        day_name = str(ws.cell(r, 2).value or "").strip()
        p2127 = float(ws.cell(r, 3).value or 0.0)
        p2128 = float(ws.cell(r, 4).value or 0.0)
        p9188 = float(ws.cell(r, 5).value or 0.0)
        notes = ws.cell(r, 7).value or ""

        if d_val:
            if hasattr(d_val, "strftime"):
                rec_date = d_val.strftime("%Y-%m-%d")
            else:
                rec_date = str(d_val)[:10]

            add_or_update_pos_record(
                record_date=rec_date,
                day_name=day_name,
                pos_2127=p2127,
                pos_2128=p2128,
                pos_9188=p9188,
                notes=str(notes),
                recorded_by="Erva Yunus"
            )
            count += 1

    print(f"[BASARILI] {count} gunluk POS cihazi kaydi sisteme aktarildi.")
    trigger_auto_backup()
    return True

if __name__ == "__main__":
    import_pos_data()
