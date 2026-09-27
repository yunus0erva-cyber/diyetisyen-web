"""
Florya MEV Koleji - Depo ve Stok Takip Sistemi
Excel İçe Aktarma (Import) ve Dışa Aktarma (Export) Araçları
"""

import io
import datetime
from typing import Dict, Any, List, Tuple
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from database import get_sqlite_connection, trigger_auto_backup
from stock_manager import add_product, record_stock_movement

TEMPLATE_COLUMNS = [
    "Barkod",
    "Ürün Kodu",
    "Ürün Adı",
    "Kategori",
    "Birim",
    "Mevcut Stok",
    "Kritik Eşik",
    "Birim Fiyat (TL)",
    "Depo Lokasyonu",
    "Son Tüketim (SKT)",
    "Parti / Lot No",
    "Açıklama / Notlar"
]

def generate_import_template() -> bytes:
    """Kullanıcının doldurması için formatlanmış hazır Excel şablonu üretir."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stok_Yukleme_Sablonu"

    # Başlık Stili
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Başlıkları yaz
    for col_num, col_name in enumerate(TEMPLATE_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_num, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Örnek 3 satır ekle
    sample_rows = [
        ["8699999001", "ORNEK-01", "Örnek Kuru Fasulye (1kg)", "Kuru Gıda & Bakliyat", "Paket", 50, 10, 85.50, "Kuru Depo", "2026-12-31", "LOT-001", "Örnek açıklama"],
        ["8699999002", "ORNEK-02", "Örnek Beyaz Peynir (500g)", "Süt ve Süt Ürünleri", "Adet", 20, 5, 120.00, "Soğuk Hava Deposu (+4°C)", "2026-10-15", "LOT-002", "Soğuk zincir"],
        ["8699999003", "ORNEK-03", "Örnek Doğal Su 1.5L", "Kantin İçecek", "Koli", 30, 10, 75.00, "Kantin Deposu", "2027-01-01", "LOT-003", "Kantin stoğu"]
    ]

    sample_font = Font(name="Segoe UI", size=10, italic=True, color="475569")
    for row_idx, row_data in enumerate(sample_rows, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = sample_font
            cell.border = thin_border

    # Sütun genişliklerini otomatik ayarla
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def parse_and_import_excel(
    file_bytes: bytes,
    update_mode: str = "update", # "update" (varsa stok artır / güncelle) veya "new_only"
    user_name: str = "Sistem"
) -> Dict[str, Any]:
    """
    Kullanıcının yüklediği Excel dosyasını doğrular ve veritabanına aktarır.
    """
    errors = []
    added_count = 0
    updated_count = 0

    try:
        df = pd.read_excel(io.BytesIO(file_bytes), dtype=str)
    except Exception as e:
        return {"success": False, "added": 0, "updated": 0, "errors": [f"Excel dosyası okunamadı: {str(e)}"]}

    # Sütun isimlerini normalize et
    df.columns = [str(c).strip() for c in df.columns]

    # Zorunlu alan kontrolü
    required_cols = ["Ürün Adı", "Kategori", "Birim"]
    for req in required_cols:
        if req not in df.columns:
            return {"success": False, "added": 0, "updated": 0, "errors": [f"Zorunlu sütun eksik: '{req}'! Lütfen örnek şablonu kullanın."]}

    conn = get_sqlite_connection()
    cursor = conn.cursor()

    for idx, row in df.iterrows():
        row_num = idx + 2 # Excel satır numarası (başlık 1)
        name = str(row.get("Ürün Adı", "")).strip()

        # Boş veya örnek satırı atla
        if not name or name == "nan" or "Örnek" in name:
            continue

        barcode = str(row.get("Barkod", "")).strip()
        if barcode == "nan" or not barcode:
            barcode = None

        product_code = str(row.get("Ürün Kodu", "")).strip()
        if product_code == "nan" or not product_code:
            product_code = None

        category = str(row.get("Kategori", "Genel")).strip()
        if category == "nan" or not category:
            category = "Genel"

        unit = str(row.get("Birim", "Adet")).strip()
        if unit == "nan" or not unit:
            unit = "Adet"

        # Sayısal değerler
        try:
            current_stock_val = float(row.get("Mevcut Stok", 0))
        except (ValueError, TypeError):
            current_stock_val = 0.0

        try:
            critical_val = float(row.get("Kritik Eşik", 5))
        except (ValueError, TypeError):
            critical_val = 5.0

        try:
            unit_price_val = float(row.get("Birim Fiyat (TL)", 0))
        except (ValueError, TypeError):
            unit_price_val = 0.0

        location = str(row.get("Depo Lokasyonu", "Kuru Depo")).strip()
        if location == "nan" or not location:
            location = "Kuru Depo"

        skt = str(row.get("Son Tüketim (SKT)", "")).strip()
        if skt == "nan" or not skt:
            skt = None
        else:
            # Tarih formatlama
            try:
                skt = pd.to_datetime(skt).strftime("%Y-%m-%d")
            except:
                skt = None

        lot_no = str(row.get("Parti / Lot No", "")).strip()
        if lot_no == "nan" or not lot_no:
            lot_no = None

        notes = str(row.get("Açıklama / Notlar", "")).strip()
        if notes == "nan" or not notes:
            notes = "Excel toplu yükleme ile eklendi"

        # Veritabanında ürün var mı kontrol et (önce barkod, sonra ürün adı)
        existing_id = None
        if barcode:
            cursor.execute("SELECT id, current_stock FROM products WHERE barcode = ?", (barcode,))
            match = cursor.fetchone()
            if match:
                existing_id = match["id"]
                old_stock = match["current_stock"]

        if not existing_id:
            cursor.execute("SELECT id, current_stock FROM products WHERE name = ?", (name,))
            match = cursor.fetchone()
            if match:
                existing_id = match["id"]
                old_stock = match["current_stock"]

        if existing_id:
            if update_mode == "update":
                # Mevcut ürüne stok ekle veya güncelle
                new_total_stock = old_stock + current_stock_val
                cursor.execute("""
                    UPDATE products
                    SET current_stock = ?,
                        critical_threshold = ?,
                        unit_price = ?,
                        storage_location = ?,
                        expiry_date = COALESCE(?, expiry_date),
                        lot_no = COALESCE(?, lot_no),
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ?
                """, (new_total_stock, critical_val, unit_price_val, location, skt, lot_no, existing_id))

                if current_stock_val > 0:
                    cursor.execute("""
                        INSERT INTO stock_movements (
                            product_id, product_name, movement_type, quantity, 
                            previous_stock, new_stock, unit, document_no, reason, created_by
                        )
                        VALUES (?, ?, 'GİRİŞ', ?, ?, ?, ?, 'EXCEL-TOPLU', 'Excel ile stok artırımı', ?)
                    """, (existing_id, name, current_stock_val, old_stock, new_total_stock, unit, user_name))

                updated_count += 1
            else:
                errors.append(f"Satır {row_num}: '{name}' ürünü zaten mevcut olduğundan atlandı.")
        else:
            # Yeni ürün olarak ekle
            success, msg = add_product(
                barcode=barcode if barcode else "",
                product_code=product_code if product_code else "",
                name=name,
                category=category,
                unit=unit,
                initial_stock=current_stock_val,
                critical_threshold=critical_val,
                unit_price=unit_price_val,
                storage_location=location,
                expiry_date=skt,
                lot_no=lot_no if lot_no else "",
                notes=notes,
                user_name=user_name
            )
            if success:
                added_count += 1
            else:
                errors.append(f"Satır {row_num} eklenemedi: {msg}")

    conn.commit()
    conn.close()
    trigger_auto_backup()

    return {
        "success": True,
        "added": added_count,
        "updated": updated_count,
        "errors": errors
    }

def export_dataframe_to_excel(
    df: pd.DataFrame,
    sheet_name: str = "Rapor",
    report_title: str = "Florya MEV Koleji Raporu",
    highlight_column: str = None
) -> bytes:
    """Pandas DataFrame'i şık ve kurumsal renklerle biçimlendirilmiş bir Excel dosyasına dönüştürür."""
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    # Başlık Başlık Satırı (Banner)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
    title_cell = ws.cell(row=1, column=1, value=f"🏫 FLORYA MEV KOLEJİ • {report_title.upper()}")
    title_cell.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Rapor Tarihi Alt Bilgisi
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(df.columns))
    now_str = datetime.datetime.now().strftime("%d.%m.%Y %H:%M")
    date_cell = ws.cell(row=2, column=1, value=f"Rapor Oluşturma Tarihi: {now_str} | Sistem: Antigravity Depo & Stok Portalı")
    date_cell.font = Font(name="Segoe UI", size=9, italic=True, color="475569")
    date_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    # Tablo Başlıkları (Row 4)
    header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=4, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 24

    # Veri Satırları (Row 5'ten itibaren)
    body_font = Font(name="Segoe UI", size=10, color="0F172A")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    # Uyarı Renkleri
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Stok yok / kritik
    yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Eşik altı

    for r_idx, row in df.iterrows():
        current_excel_row = r_idx + 5
        is_even = (r_idx % 2 == 0)
        row_fill = zebra_fill if is_even else None

        # Eğer kritik uyarı satırıysa arka planı renklendir
        if "Durum Seviyesi" in df.columns:
            status_val = str(row.get("Durum Seviyesi", ""))
            if "TÜKENDİ" in status_val or "ÇOK ACİL" in status_val:
                row_fill = red_fill
            elif "DİKKAT" in status_val:
                row_fill = yellow_fill

        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=current_excel_row, column=c_idx)
            
            # Sayı veya tarih formatı kontrolü
            if pd.isna(val) or val is None:
                cell.value = "-"
            elif isinstance(val, (int, float)):
                cell.value = round(val, 2)
                col_title = df.columns[c_idx - 1]
                if "Fiyat" in col_title or "Tutar" in col_title or "Değer" in col_title:
                    cell.number_format = '#,##0.00 "₺"'
                else:
                    cell.number_format = '#,##0.00'
            else:
                cell.value = str(val)

            cell.font = body_font
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

        ws.row_dimensions[current_excel_row].height = 20

    # Sütun genişliklerini ayarla
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row > 2: # Başlık bannerlarını hesaplamaya katma
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output)
    return output.getvalue()
