# -*- coding: utf-8 -*-
"""
Florya MEV Koleji & NOYA Kantin Besin Değerleri Tablosu Oluşturucu
Google E-Tablolar ve Excel Uyumlu Rapor
"""

import os
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ürün Veri Seti (57 Ürün)
DATA = [
    # YİYECEK GRUBU
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "POĞAÇA",
        "fiyat": 35.0,
        "porsiyon": "1 adet (~80 g)",
        "gramaj_g": 80,
        "kalori_kcal": 305,
        "protein_g": 6.8,
        "karbonhidrat_g": 34.2,
        "yag_g": 16.0,
        "seker_g": 2.5,
        "doymus_yag_g": 5.5,
        "lif_g": 1.8,
        "tuz_g": 1.2,
        "aciklama": "Sade / Peynirli fırın poğaça"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "AÇMA",
        "fiyat": 35.0,
        "porsiyon": "1 adet (~105 g)",
        "gramaj_g": 105,
        "kalori_kcal": 378,
        "protein_g": 7.6,
        "karbonhidrat_g": 48.5,
        "yag_g": 17.2,
        "seker_g": 4.8,
        "doymus_yag_g": 6.0,
        "lif_g": 2.1,
        "tuz_g": 1.3,
        "aciklama": "Sade yumuşak fırın açması"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "SİMİT",
        "fiyat": 35.0,
        "porsiyon": "1 adet (~100 g)",
        "gramaj_g": 100,
        "kalori_kcal": 302,
        "protein_g": 10.2,
        "karbonhidrat_g": 57.4,
        "yag_g": 4.1,
        "seker_g": 3.2,
        "doymus_yag_g": 0.8,
        "lif_g": 3.5,
        "tuz_g": 1.5,
        "aciklama": "Geleneksel susamlı sokak simidi"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "KAŞARLI TOST",
        "fiyat": 90.0,
        "porsiyon": "1 porsiyon (~125 g)",
        "gramaj_g": 125,
        "kalori_kcal": 352,
        "protein_g": 14.8,
        "karbonhidrat_g": 33.5,
        "yag_g": 18.2,
        "seker_g": 2.1,
        "doymus_yag_g": 9.8,
        "lif_g": 1.6,
        "tuz_g": 1.8,
        "aciklama": "2 dilim tost ekmeği, 50g taze kaşar, tereyağı"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "AYVALIK TOST",
        "fiyat": 120.0,
        "porsiyon": "1 porsiyon (~270 g)",
        "gramaj_g": 270,
        "kalori_kcal": 645,
        "protein_g": 25.4,
        "karbonhidrat_g": 58.2,
        "yag_g": 35.6,
        "seker_g": 4.5,
        "doymus_yag_g": 14.2,
        "lif_g": 2.8,
        "tuz_g": 3.2,
        "aciklama": "Ayvalık ekmeği, sucuk, sosis, salam, kaşar, turşu, sos"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "PİTOS",
        "fiyat": 120.0,
        "porsiyon": "1 porsiyon (~190 g)",
        "gramaj_g": 190,
        "kalori_kcal": 508,
        "protein_g": 19.5,
        "karbonhidrat_g": 52.0,
        "yag_g": 24.8,
        "seker_g": 3.8,
        "doymus_yag_g": 11.2,
        "lif_g": 2.4,
        "tuz_g": 2.6,
        "aciklama": "Pide ekmeğinde kaşar, sucuk ve domates soslu sıcak tost"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "KUMRU",
        "fiyat": 120.0,
        "porsiyon": "1 porsiyon (~230 g)",
        "gramaj_g": 230,
        "kalori_kcal": 595,
        "protein_g": 22.8,
        "karbonhidrat_g": 54.5,
        "yag_g": 32.4,
        "seker_g": 4.2,
        "doymus_yag_g": 12.8,
        "lif_g": 2.6,
        "tuz_g": 3.0,
        "aciklama": "Susamlı kumru ekmeği, ızgara sucuk, salam, sosis, kaşar, domates"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "BİOFRESH SAND. VE ÜÇ PEYNİRLİ BAGEL",
        "fiyat": 120.0,
        "porsiyon": "1 paket (~170 g)",
        "gramaj_g": 170,
        "kalori_kcal": 412,
        "protein_g": 16.4,
        "karbonhidrat_g": 46.8,
        "yag_g": 17.5,
        "seker_g": 3.6,
        "doymus_yag_g": 8.4,
        "lif_g": 2.5,
        "tuz_g": 2.1,
        "aciklama": "Biofresh soğuk sandviç / Bagel, beyaz peynir, kaşar, labne"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "ÇİĞKÖFTE DÜRÜM",
        "fiyat": 130.0,
        "porsiyon": "1 porsiyon (~200 g)",
        "gramaj_g": 200,
        "kalori_kcal": 385,
        "protein_g": 9.4,
        "karbonhidrat_g": 63.8,
        "yag_g": 10.2,
        "seker_g": 6.5,
        "doymus_yag_g": 1.4,
        "lif_g": 6.2,
        "tuz_g": 2.4,
        "aciklama": "Lavaş ekmeği, etsiz çiğköfte, kıvırcık marul, nar ekşisi, limon"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "IZGARA KÖFTE SANDVİÇ",
        "fiyat": 170.0,
        "porsiyon": "1 porsiyon (~240 g)",
        "gramaj_g": 240,
        "kalori_kcal": 538,
        "protein_g": 27.6,
        "karbonhidrat_g": 51.5,
        "yag_g": 24.5,
        "seker_g": 3.5,
        "doymus_yag_g": 9.2,
        "lif_g": 3.1,
        "tuz_g": 2.8,
        "aciklama": "Sandviç ekmeği, 4 adet ızgara dana köfte, domates, marul"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "EKMEK ARASI ŞİNİTZEL",
        "fiyat": 160.0,
        "porsiyon": "1 porsiyon (~240 g)",
        "gramaj_g": 240,
        "kalori_kcal": 565,
        "protein_g": 24.2,
        "karbonhidrat_g": 58.6,
        "yag_g": 26.2,
        "seker_g": 3.2,
        "doymus_yag_g": 5.6,
        "lif_g": 2.8,
        "tuz_g": 2.5,
        "aciklama": "Sandviç ekmeği, çıtır panelenmiş tavuk şinitzel, yeşillik, sos"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "PİZZA",
        "fiyat": 100.0,
        "porsiyon": "1 dilim (~160 g)",
        "gramaj_g": 160,
        "kalori_kcal": 408,
        "protein_g": 14.5,
        "karbonhidrat_g": 47.2,
        "yag_g": 18.0,
        "seker_g": 3.8,
        "doymus_yag_g": 7.2,
        "lif_g": 2.7,
        "tuz_g": 2.2,
        "aciklama": "Kantin karışık pizza (kaşar, sucuk, mısır, biber, domates sos)"
    },
    {
        "kategori": "YİYECEK GRUBU",
        "urun": "TAZE MEYVE",
        "fiyat": 40.0,
        "porsiyon": "1 adet (~150 g)",
        "gramaj_g": 150,
        "kalori_kcal": 82,
        "protein_g": 0.9,
        "karbonhidrat_g": 20.4,
        "yag_g": 0.3,
        "seker_g": 16.5,
        "doymus_yag_g": 0.1,
        "lif_g": 3.2,
        "tuz_g": 0.01,
        "aciklama": "Mevsim meyvesi (Elma / Muz / Mandalina ortalaması)"
    },

    # TATLI GRUBU
    {
        "kategori": "TATLI GRUBU",
        "urun": "FRAMBUAZLU CUP",
        "fiyat": 130.0,
        "porsiyon": "1 kupa (~150 g)",
        "gramaj_g": 150,
        "kalori_kcal": 325,
        "protein_g": 4.2,
        "karbonhidrat_g": 43.5,
        "yag_g": 15.0,
        "seker_g": 28.4,
        "doymus_yag_g": 8.5,
        "lif_g": 1.8,
        "tuz_g": 0.3,
        "aciklama": "Bisküvi tabanı, hafif krema, doğal frambuaz püresi"
    },
    {
        "kategori": "TATLI GRUBU",
        "urun": "ÇİKOLATALI CUP",
        "fiyat": 130.0,
        "porsiyon": "1 kupa (~150 g)",
        "gramaj_g": 150,
        "kalori_kcal": 368,
        "protein_g": 5.4,
        "karbonhidrat_g": 46.2,
        "yag_g": 18.5,
        "seker_g": 32.6,
        "doymus_yag_g": 10.8,
        "lif_g": 2.2,
        "tuz_g": 0.4,
        "aciklama": "Kakaolu kek, çikolata mus krema, parça çikolata"
    },
    {
        "kategori": "TATLI GRUBU",
        "urun": "LOTUSLU CUP",
        "fiyat": 130.0,
        "porsiyon": "1 kupa (~150 g)",
        "gramaj_g": 150,
        "kalori_kcal": 395,
        "protein_g": 4.5,
        "karbonhidrat_g": 49.0,
        "yag_g": 20.4,
        "seker_g": 34.8,
        "doymus_yag_g": 10.2,
        "lif_g": 1.2,
        "tuz_g": 0.5,
        "aciklama": "Lotus Biscoff bisküvi kırıntısı, karamelize bisküvi kreması"
    },
    {
        "kategori": "TATLI GRUBU",
        "urun": "TİRAMİSU CUP",
        "fiyat": 170.0,
        "porsiyon": "1 kupa (~150 g)",
        "gramaj_g": 150,
        "kalori_kcal": 342,
        "protein_g": 5.2,
        "karbonhidrat_g": 39.4,
        "yag_g": 18.2,
        "seker_g": 26.5,
        "doymus_yag_g": 10.5,
        "lif_g": 1.4,
        "tuz_g": 0.3,
        "aciklama": "Kedidili bisküvi, kahve şurubu, mascarpone/labne krema, kakao"
    },
    {
        "kategori": "TATLI GRUBU",
        "urun": "DONUT SİYAH",
        "fiyat": 80.0,
        "porsiyon": "1 adet (~65 g)",
        "gramaj_g": 65,
        "kalori_kcal": 275,
        "protein_g": 4.1,
        "karbonhidrat_g": 33.8,
        "yag_g": 14.2,
        "seker_g": 16.5,
        "doymus_yag_g": 6.8,
        "lif_g": 1.5,
        "tuz_g": 0.4,
        "aciklama": "Çikolata kaplamalı ve kakaolu süslemeli donut"
    },
    {
        "kategori": "TATLI GRUBU",
        "urun": "DONUT PEMBE",
        "fiyat": 80.0,
        "porsiyon": "1 adet (~65 g)",
        "gramaj_g": 65,
        "kalori_kcal": 265,
        "protein_g": 3.8,
        "karbonhidrat_g": 35.2,
        "yag_g": 13.0,
        "seker_g": 17.8,
        "doymus_yag_g": 6.2,
        "lif_g": 1.2,
        "tuz_g": 0.4,
        "aciklama": "Çilek aromalı pembe glazür kaplamalı donut"
    },

    # DONDURMA GRUBU
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "MAGNUM",
        "fiyat": 100.0,
        "porsiyon": "1 adet (~73 g / 100 ml)",
        "gramaj_g": 73,
        "kalori_kcal": 257,
        "protein_g": 3.2,
        "karbonhidrat_g": 24.0,
        "yag_g": 16.0,
        "seker_g": 22.0,
        "doymus_yag_g": 11.0,
        "lif_g": 0.8,
        "tuz_g": 0.1,
        "aciklama": "Magnum Classic vanilyalı dondurma ve kalın sütlü çikolata kaplama"
    },
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "CORNETTO DISC",
        "fiyat": 90.0,
        "porsiyon": "1 adet (~87 g / 140 ml)",
        "gramaj_g": 87,
        "kalori_kcal": 285,
        "protein_g": 3.6,
        "karbonhidrat_g": 35.4,
        "yag_g": 14.2,
        "seker_g": 25.2,
        "doymus_yag_g": 8.6,
        "lif_g": 1.1,
        "tuz_g": 0.2,
        "aciklama": "Gevrek külah, vanilyalı krema ve üst çikolata diski"
    },
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "NOGGER",
        "fiyat": 75.0,
        "porsiyon": "1 adet (~75 g / 145 ml)",
        "gramaj_g": 75,
        "kalori_kcal": 222,
        "protein_g": 3.9,
        "karbonhidrat_g": 30.2,
        "yag_g": 9.6,
        "seker_g": 18.5,
        "doymus_yag_g": 5.4,
        "lif_g": 1.2,
        "tuz_g": 0.2,
        "aciklama": "Nogger Sandwich bisküvili karamelli dondurma"
    },
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "VOLCANO",
        "fiyat": 60.0,
        "porsiyon": "1 adet (~60 g / 95 ml)",
        "gramaj_g": 60,
        "kalori_kcal": 176,
        "protein_g": 2.3,
        "karbonhidrat_g": 22.4,
        "yag_g": 8.8,
        "seker_g": 16.8,
        "doymus_yag_g": 5.2,
        "lif_g": 0.6,
        "tuz_g": 0.1,
        "aciklama": "Külah içi çikolata sos volkan dolgulu dondurma"
    },
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "BOOM BOOM",
        "fiyat": 60.0,
        "porsiyon": "1 adet (~55 g / 85 ml)",
        "gramaj_g": 55,
        "kalori_kcal": 156,
        "protein_g": 1.9,
        "karbonhidrat_g": 19.2,
        "yag_g": 8.1,
        "seker_g": 15.4,
        "doymus_yag_g": 5.0,
        "lif_g": 0.5,
        "tuz_g": 0.1,
        "aciklama": "Çikolata kaplamalı çubuk dondurma"
    },
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "FRIGOLA",
        "fiyat": 45.0,
        "porsiyon": "1 adet (~45 g / 60 ml)",
        "gramaj_g": 45,
        "kalori_kcal": 148,
        "protein_g": 2.1,
        "karbonhidrat_g": 16.5,
        "yag_g": 8.4,
        "seker_g": 14.2,
        "doymus_yag_g": 5.1,
        "lif_g": 0.7,
        "tuz_g": 0.1,
        "aciklama": "Kakaolu bisküvi parçalı klasik Frigola çubuk dondurma"
    },
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "CLASSICS",
        "fiyat": 30.0,
        "porsiyon": "1 adet (~45 g / 70 ml)",
        "gramaj_g": 45,
        "kalori_kcal": 125,
        "protein_g": 1.6,
        "karbonhidrat_g": 14.5,
        "yag_g": 6.9,
        "seker_g": 12.8,
        "doymus_yag_g": 4.4,
        "lif_g": 0.4,
        "tuz_g": 0.1,
        "aciklama": "Algida Classics sütlü çikolata kaplı çubuk dondurma"
    },
    {
        "kategori": "DONDURMA GRUBU",
        "urun": "TWISTER",
        "fiyat": 20.0,
        "porsiyon": "1 adet (~62 g / 70 ml)",
        "gramaj_g": 62,
        "kalori_kcal": 68,
        "protein_g": 0.5,
        "karbonhidrat_g": 16.2,
        "yag_g": 0.8,
        "seker_g": 13.5,
        "doymus_yag_g": 0.5,
        "lif_g": 0.3,
        "tuz_g": 0.03,
        "aciklama": "Meyve aromalı spiral su/süt buzu dondurma"
    },

    # ETİ GRUBU
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ BENİMO LOKMALIK",
        "fiyat": 40.0,
        "porsiyon": "1 paket (72 g)",
        "gramaj_g": 72,
        "kalori_kcal": 338,
        "protein_g": 4.2,
        "karbonhidrat_g": 51.8,
        "yag_g": 12.2,
        "seker_g": 32.4,
        "doymus_yag_g": 7.2,
        "lif_g": 1.8,
        "tuz_g": 0.4,
        "aciklama": "Marshmallowlu hindistan cevizli lokmalık bisküvi"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ BİDOLU",
        "fiyat": 30.0,
        "porsiyon": "1 paket (36 g)",
        "gramaj_g": 36,
        "kalori_kcal": 193,
        "protein_g": 3.3,
        "karbonhidrat_g": 19.4,
        "yag_g": 11.2,
        "seker_g": 12.6,
        "doymus_yag_g": 5.8,
        "lif_g": 1.1,
        "tuz_g": 0.2,
        "aciklama": "Fıstık ezmeli krema dolgulu çikolata kaplı gofret"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ BROWNIE GOLD KEK",
        "fiyat": 30.0,
        "porsiyon": "1 paket (45 g)",
        "gramaj_g": 45,
        "kalori_kcal": 205,
        "protein_g": 2.6,
        "karbonhidrat_g": 23.4,
        "yag_g": 11.4,
        "seker_g": 17.1,
        "doymus_yag_g": 5.2,
        "lif_g": 1.4,
        "tuz_g": 0.3,
        "aciklama": "Çikolata soslu ve fındıklı ıslak kek"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ BROWNIE INTENSE",
        "fiyat": 40.0,
        "porsiyon": "1 paket (50 g)",
        "gramaj_g": 50,
        "kalori_kcal": 234,
        "protein_g": 2.8,
        "karbonhidrat_g": 27.0,
        "yag_g": 13.0,
        "seker_g": 20.5,
        "doymus_yag_g": 6.8,
        "lif_g": 1.6,
        "tuz_g": 0.3,
        "aciklama": "Krema dolgulu, çikolata kaplamalı yoğun brownie kek"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ CANGA",
        "fiyat": 40.0,
        "porsiyon": "1 paket (45 g)",
        "gramaj_g": 45,
        "kalori_kcal": 232,
        "protein_g": 4.6,
        "karbonhidrat_g": 24.3,
        "yag_g": 13.1,
        "seker_g": 19.4,
        "doymus_yag_g": 5.4,
        "lif_g": 1.8,
        "tuz_g": 0.2,
        "aciklama": "Bol yer fıstıklı, karamelli nuga sütlü çikolata"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ CİN LOKMALIK",
        "fiyat": 40.0,
        "porsiyon": "1 paket (114 g)",
        "gramaj_g": 114,
        "kalori_kcal": 482,
        "protein_g": 5.2,
        "karbonhidrat_g": 84.4,
        "yag_g": 13.5,
        "seker_g": 48.2,
        "doymus_yag_g": 6.8,
        "lif_g": 2.4,
        "tuz_g": 0.5,
        "aciklama": "Portakal pelteli ve renkli granüllü lokmalık bisküvi"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ CRAX ÇEŞİT",
        "fiyat": 30.0,
        "porsiyon": "1 paket (50 g)",
        "gramaj_g": 50,
        "kalori_kcal": 204,
        "protein_g": 4.6,
        "karbonhidrat_g": 33.2,
        "yag_g": 6.0,
        "seker_g": 2.1,
        "doymus_yag_g": 2.6,
        "lif_g": 1.8,
        "tuz_g": 1.4,
        "aciklama": "Sade / Baharatlı fırınlanmış çubuk kraker"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ CRAX PATLAYAN LEZZET",
        "fiyat": 40.0,
        "porsiyon": "1 paket (50 g)",
        "gramaj_g": 50,
        "kalori_kcal": 212,
        "protein_g": 4.3,
        "karbonhidrat_g": 32.5,
        "yag_g": 7.2,
        "seker_g": 2.4,
        "doymus_yag_g": 3.1,
        "lif_g": 1.7,
        "tuz_g": 1.5,
        "aciklama": "Ekstra baharat ve aroma kaplı patlayan lezzet kraker"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ GONG",
        "fiyat": 40.0,
        "porsiyon": "1 paket (64 g)",
        "gramaj_g": 64,
        "kalori_kcal": 284,
        "protein_g": 4.9,
        "karbonhidrat_g": 46.4,
        "yag_g": 8.6,
        "seker_g": 3.8,
        "doymus_yag_g": 3.8,
        "lif_g": 2.2,
        "tuz_g": 1.3,
        "aciklama": "Fırınlanmış çıtır mısır ve pirinç patlağı"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ HOŞBEŞ",
        "fiyat": 40.0,
        "porsiyon": "1 paket (66 g)",
        "gramaj_g": 66,
        "kalori_kcal": 339,
        "protein_g": 4.1,
        "karbonhidrat_g": 42.2,
        "yag_g": 17.2,
        "seker_g": 24.2,
        "doymus_yag_g": 9.2,
        "lif_g": 1.5,
        "tuz_g": 0.3,
        "aciklama": "İncecik çıtır yaprak gofret (Fındıklı/Çilekli)"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ KARAM GURME",
        "fiyat": 40.0,
        "porsiyon": "1 paket (50 g)",
        "gramaj_g": 50,
        "kalori_kcal": 268,
        "protein_g": 3.6,
        "karbonhidrat_g": 28.2,
        "yag_g": 15.6,
        "seker_g": 20.4,
        "doymus_yag_g": 8.5,
        "lif_g": 2.6,
        "tuz_g": 0.2,
        "aciklama": "Yoğun bitter çikolata kremalı gofret"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ KOMBO",
        "fiyat": 40.0,
        "porsiyon": "1 paket (56 g)",
        "gramaj_g": 56,
        "kalori_kcal": 258,
        "protein_g": 3.2,
        "karbonhidrat_g": 37.2,
        "yag_g": 10.6,
        "seker_g": 24.5,
        "doymus_yag_g": 5.8,
        "lif_g": 1.4,
        "tuz_g": 0.3,
        "aciklama": "Bisküvi, marshmallow ve sütlü çikolata kombinasyonu"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ NERO",
        "fiyat": 40.0,
        "porsiyon": "1 paket (100 g)",
        "gramaj_g": 100,
        "kalori_kcal": 486,
        "protein_g": 5.9,
        "karbonhidrat_g": 66.2,
        "yag_g": 21.4,
        "seker_g": 32.5,
        "doymus_yag_g": 10.5,
        "lif_g": 3.4,
        "tuz_g": 0.7,
        "aciklama": "Kakaolu bisküvi arası vanilyalı krema dolgulu bisküvi"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ PATİTO PATİ",
        "fiyat": 25.0,
        "porsiyon": "1 paket (35 g)",
        "gramaj_g": 35,
        "kalori_kcal": 184,
        "protein_g": 2.2,
        "karbonhidrat_g": 19.2,
        "yag_g": 11.2,
        "seker_g": 1.1,
        "doymus_yag_g": 4.8,
        "lif_g": 1.2,
        "tuz_g": 0.6,
        "aciklama": "Tuzlu çıtır patates cips atıştırmalık"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ PİZZA KRAKER",
        "fiyat": 30.0,
        "porsiyon": "1 paket (45 g)",
        "gramaj_g": 45,
        "kalori_kcal": 200,
        "protein_g": 3.9,
        "karbonhidrat_g": 28.5,
        "yag_g": 7.8,
        "seker_g": 2.5,
        "doymus_yag_g": 3.5,
        "lif_g": 1.6,
        "tuz_g": 1.2,
        "aciklama": "Pizza baharatlı çıtır kraker"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ POPKEK",
        "fiyat": 30.0,
        "porsiyon": "1 paket (60 g)",
        "gramaj_g": 60,
        "kalori_kcal": 266,
        "protein_g": 3.1,
        "karbonhidrat_g": 34.5,
        "yag_g": 12.8,
        "seker_g": 22.4,
        "doymus_yag_g": 6.2,
        "lif_g": 1.5,
        "tuz_g": 0.4,
        "aciklama": "Sütlü çikolata kaplı, muz / limon kremalı kek"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ SUSAMLI ÇUBUK",
        "fiyat": 30.0,
        "porsiyon": "1 paket (60 g)",
        "gramaj_g": 60,
        "kalori_kcal": 278,
        "protein_g": 6.2,
        "karbonhidrat_g": 38.4,
        "yag_g": 11.2,
        "seker_g": 1.6,
        "doymus_yag_g": 2.2,
        "lif_g": 2.5,
        "tuz_g": 1.4,
        "aciklama": "Kavrulmuş bol susam kaplı gevrek çubuk kraker"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ SÜT BURGER",
        "fiyat": 30.0,
        "porsiyon": "1 paket (35 g)",
        "gramaj_g": 35,
        "kalori_kcal": 145,
        "protein_g": 2.4,
        "karbonhidrat_g": 17.6,
        "yag_g": 7.3,
        "seker_g": 13.2,
        "doymus_yag_g": 4.2,
        "lif_g": 0.6,
        "tuz_g": 0.2,
        "aciklama": "İki yumuşak kek arası ballı süt kremalı soğuk sandviç"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ TOPKEK",
        "fiyat": 30.0,
        "porsiyon": "1 paket (40 g)",
        "gramaj_g": 40,
        "kalori_kcal": 173,
        "protein_g": 2.2,
        "karbonhidrat_g": 23.2,
        "yag_g": 7.9,
        "seker_g": 14.1,
        "doymus_yag_g": 3.6,
        "lif_g": 0.9,
        "tuz_g": 0.3,
        "aciklama": "Portakallı / Kakaolu yumuşak mini kek"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ TUTKU",
        "fiyat": 40.0,
        "porsiyon": "1 paket (100 g)",
        "gramaj_g": 100,
        "kalori_kcal": 508,
        "protein_g": 4.8,
        "karbonhidrat_g": 64.5,
        "yag_g": 25.8,
        "seker_g": 31.2,
        "doymus_yag_g": 12.4,
        "lif_g": 2.8,
        "tuz_g": 0.5,
        "aciklama": "Mozaik desenli akışkan kakaolu krema dolgulu bisküvi"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ UZUN STİCK-A.FISTIKLI",
        "fiyat": 60.0,
        "porsiyon": "1 paket (34 g)",
        "gramaj_g": 34,
        "kalori_kcal": 185,
        "protein_g": 3.3,
        "karbonhidrat_g": 17.2,
        "yag_g": 11.6,
        "seker_g": 15.2,
        "doymus_yag_g": 5.6,
        "lif_g": 1.2,
        "tuz_g": 0.1,
        "aciklama": "Bütün Antep fıstıklı sütlü çikolata çubuğu"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ UZUN STİCK-SÜTLÜ",
        "fiyat": 50.0,
        "porsiyon": "1 paket (34 g)",
        "gramaj_g": 34,
        "kalori_kcal": 181,
        "protein_g": 2.6,
        "karbonhidrat_g": 19.2,
        "yag_g": 10.6,
        "seker_g": 18.2,
        "doymus_yag_g": 6.4,
        "lif_g": 0.8,
        "tuz_g": 0.1,
        "aciklama": "Saf sütlü çikolata stick bar"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ BALIK KRAKER MISIRLI",
        "fiyat": 30.0,
        "porsiyon": "1 paket (40 g)",
        "gramaj_g": 40,
        "kalori_kcal": 175,
        "protein_g": 3.5,
        "karbonhidrat_g": 27.2,
        "yag_g": 5.8,
        "seker_g": 2.1,
        "doymus_yag_g": 2.5,
        "lif_g": 1.4,
        "tuz_g": 1.1,
        "aciklama": "Mısır çeşnili gevrek balık figürlü kraker"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ LİFALİF BAR",
        "fiyat": 70.0,
        "porsiyon": "1 paket (35 g)",
        "gramaj_g": 35,
        "kalori_kcal": 142,
        "protein_g": 3.6,
        "karbonhidrat_g": 19.5,
        "yag_g": 5.3,
        "seker_g": 8.2,
        "doymus_yag_g": 1.1,
        "lif_g": 3.4,
        "tuz_g": 0.1,
        "aciklama": "İlave şekersiz, yulaf ve kuru yemişli tahıl bar"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ BURÇAK TUZLU-ÇÖREKOTLU",
        "fiyat": 40.0,
        "porsiyon": "1 paket (91 g)",
        "gramaj_g": 91,
        "kalori_kcal": 421,
        "protein_g": 6.9,
        "karbonhidrat_g": 55.4,
        "yag_g": 18.4,
        "seker_g": 9.2,
        "doymus_yag_g": 7.8,
        "lif_g": 5.2,
        "tuz_g": 1.6,
        "aciklama": "Tam buğday unlu, çörekotlu ve susamlı tuzlu bisküvi"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ FORM",
        "fiyat": 40.0,
        "porsiyon": "1 paket (45 g)",
        "gramaj_g": 45,
        "kalori_kcal": 168,
        "protein_g": 4.1,
        "karbonhidrat_g": 30.4,
        "yag_g": 3.6,
        "seker_g": 2.2,
        "doymus_yag_g": 0.9,
        "lif_g": 4.6,
        "tuz_g": 0.8,
        "aciklama": "Yüksek lifli, kepekli ve düşük yağlı diyet kraker"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ PUF",
        "fiyat": 20.0,
        "porsiyon": "1 paket (18 g)",
        "gramaj_g": 18,
        "kalori_kcal": 75,
        "protein_g": 0.8,
        "karbonhidrat_g": 14.2,
        "yag_g": 1.7,
        "seker_g": 9.6,
        "doymus_yag_g": 0.8,
        "lif_g": 0.4,
        "tuz_g": 0.1,
        "aciklama": "Bisküvi üzeri marshmallow ve renkli pasta granülü"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ WANTED KARAMELLİ",
        "fiyat": 30.0,
        "porsiyon": "1 paket (32 g)",
        "gramaj_g": 32,
        "kalori_kcal": 158,
        "protein_g": 1.9,
        "karbonhidrat_g": 20.2,
        "yag_g": 7.8,
        "seker_g": 16.2,
        "doymus_yag_g": 4.2,
        "lif_g": 0.7,
        "tuz_g": 0.2,
        "aciklama": "Pirinç patlaklı, karamelli sütlü çikolata bar"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ WANTED BUMBA",
        "fiyat": 30.0,
        "porsiyon": "1 paket (32 g)",
        "gramaj_g": 32,
        "kalori_kcal": 162,
        "protein_g": 2.1,
        "karbonhidrat_g": 19.4,
        "yag_g": 8.4,
        "seker_g": 15.5,
        "doymus_yag_g": 4.5,
        "lif_g": 0.9,
        "tuz_g": 0.2,
        "aciklama": "Fındıklı ve kakaolu pirinç patlaklı çikolata bar"
    },
    {
        "kategori": "ETİ GRUBU",
        "urun": "ETİ MAXIMUS",
        "fiyat": 40.0,
        "porsiyon": "1 paket (36 g)",
        "gramaj_g": 36,
        "kalori_kcal": 184,
        "protein_g": 3.6,
        "karbonhidrat_g": 20.4,
        "yag_g": 9.9,
        "seker_g": 15.4,
        "doymus_yag_g": 4.6,
        "lif_g": 1.2,
        "tuz_g": 0.2,
        "aciklama": "Bol yer fıstıklı, karamel nuga ve sütlü çikolata bar"
    }
]

def create_nutrition_workbook():
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # 1. SAYFA: TÜM ÜRÜNLER VE BESİN DEĞERLERİ (MASTER DATA)
    # -------------------------------------------------------------
    ws_all = wb.active
    ws_all.title = "Kantin Besin Değerleri"
    ws_all.views.sheetView[0].showGridLines = True
    
    # Renk Paleti (Florya MEV / Noya Kurumsal Tasarım)
    navy_dark = "1B365D"      # Başlık laciverti
    navy_light = "2E5B88"     # Alt başlık laciverti
    gray_header = "4A5568"    # Koyu gri
    zebra_even = "F8FAFC"     # Çok açık buz mavisi/gri
    zebra_odd = "FFFFFF"      # Beyaz
    border_gray = "CBD5E1"    # Kenarlık rengi
    
    cat_colors = {
        "YİYECEK GRUBU": "EBF8FF",   # Açık mavi
        "TATLI GRUBU": "FAF5FF",     # Açık leylak
        "DONDURMA GRUBU": "F0FDF4",  # Açık nane yeşili
        "ETİ GRUBU": "FFFBEB"        # Açık bal sarısı
    }
    
    # Başlık Bloğu
    ws_all.merge_cells("A1:M1")
    title_cell = ws_all["A1"]
    title_cell.value = "FLORYA MEV KOLEJİ - NOYA KANTİN BESİN DEĞERLERİ VE KALORİ CETVELİ (2026-2027)"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_all.row_dimensions[1].height = 36

    ws_all.merge_cells("A2:M2")
    sub_cell = ws_all["A2"]
    sub_cell.value = "Tüm Ürünlerin Porsiyon Gramajı, Kalori (kcal), Protein, Karbonhidrat, Yağ, Şeker, Lif ve Tuz Analiz Tablosu"
    sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="FFFFFF")
    sub_cell.fill = PatternFill(start_color=navy_light, end_color=navy_light, fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_all.row_dimensions[2].height = 22

    # Tablo Başlıkları
    headers = [
        "Kategori",
        "Ürün Adı",
        "Satış Fiyatı (₺)",
        "Porsiyon / Gramaj",
        "Gram (g)",
        "Kalori (kcal)",
        "Protein (g)",
        "Karbonhidrat (g)",
        "Yağ (g)",
        "Şeker (g)",
        "Doymuş Yağ (g)",
        "Lif (g)",
        "Tuz (g)",
        "Ürün İçeriği ve Açıklama"
    ]
    
    # Yeniden merge gerekirse A1:N1 yapalım çünkü 14 kolon oldu
    ws_all.unmerge_cells("A1:M1")
    ws_all.unmerge_cells("A2:M2")
    ws_all.merge_cells("A1:N1")
    ws_all.merge_cells("A2:N2")
    
    header_row = 4
    ws_all.row_dimensions[header_row].height = 28
    
    thin_border = Border(
        left=Side(style='thin', color=border_gray),
        right=Side(style='thin', color=border_gray),
        top=Side(style='thin', color=border_gray),
        bottom=Side(style='thin', color=border_gray)
    )
    
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws_all.cell(row=header_row, column=col_idx, value=h_text)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=gray_header, end_color=gray_header, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Satırları Ekle
    current_row = 5
    for item in DATA:
        is_even = (current_row % 2 == 0)
        row_fill_color = zebra_even if is_even else zebra_odd
        
        row_values = [
            item["kategori"],
            item["urun"],
            item["fiyat"],
            item["porsiyon"],
            item["gramaj_g"],
            item["kalori_kcal"],
            item["protein_g"],
            item["karbonhidrat_g"],
            item["yag_g"],
            item["seker_g"],
            item["doymus_yag_g"],
            item["lif_g"],
            item["tuz_g"],
            item["aciklama"]
        ]
        
        ws_all.row_dimensions[current_row].height = 22
        for col_idx, val in enumerate(row_values, 1):
            cell = ws_all.cell(row=current_row, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=9.5)
            cell.border = thin_border
            cell.fill = PatternFill(start_color=row_fill_color, end_color=row_fill_color, fill_type="solid")
            
            # Hizalama ve Sayı Formatları
            if col_idx in [1]:  # Kategori
                cell.alignment = Alignment(horizontal="center", vertical="center")
                # Kategoriye hafif özel renk tonu
                cat_col = cat_colors.get(val, row_fill_color)
                cell.fill = PatternFill(start_color=cat_col, end_color=cat_col, fill_type="solid")
                cell.font = Font(name="Segoe UI", size=9.5, bold=True)
            elif col_idx in [2]: # Ürün Adı
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = Font(name="Segoe UI", size=9.5, bold=True)
            elif col_idx == 3: # Fiyat
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '₺#,##0.00'
            elif col_idx == 4: # Porsiyon
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [5, 6]: # Gramaj, Kalori
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
                if col_idx == 6:
                    cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="C026D3")
            elif col_idx in [7, 8, 9, 10, 11, 12, 13]: # Besin Değerleri (g)
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.0'
            else: # Açıklama
                cell.alignment = Alignment(horizontal="left", vertical="center")
                
        current_row += 1

    # Alt Toplam / Ortalama Satırı
    total_row = current_row
    ws_all.row_dimensions[total_row].height = 25
    ws_all.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=2)
    avg_cell = ws_all.cell(row=total_row, column=1, value="GENEL ORTALAMA (Tüm Ürünler)")
    avg_cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    avg_cell.alignment = Alignment(horizontal="center", vertical="center")
    avg_cell.fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    ws_all.cell(row=total_row, column=2).fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")

    for col_idx in range(3, 15):
        cell = ws_all.cell(row=total_row, column=col_idx)
        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
        cell.border = thin_border
        
        col_letter = get_column_letter(col_idx)
        if col_idx == 3: # Ortalama Fiyat
            cell.value = f"=AVERAGE({col_letter}5:{col_letter}{total_row-1})"
            cell.number_format = '₺#,##0.00'
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx == 4:
            cell.value = "-"
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [5, 6]:
            cell.value = f"=AVERAGE({col_letter}5:{col_letter}{total_row-1})"
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx in [7, 8, 9, 10, 11, 12, 13]:
            cell.value = f"=AVERAGE({col_letter}5:{col_letter}{total_row-1})"
            cell.number_format = '#,##0.0'
            cell.alignment = Alignment(horizontal="right", vertical="center")
        elif col_idx == 14:
            cell.value = "57 Farklı Ürün Analizi"
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Kolon Genişliklerini Ayarla
    for col in ws_all.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_all.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------
    # 2. SAYFA: KATEGORİ ÖZETİ VE BESLENME ANALİZİ
    # -------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Kategori Özeti")
    ws_summary.views.sheetView[0].showGridLines = True
    
    ws_summary.merge_cells("A1:H1")
    s_title = ws_summary["A1"]
    s_title.value = "KATEGORİ BAZLI BESİN VE KALORİ ÖZET RAPORU"
    s_title.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    s_title.fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    s_title.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 32

    s_headers = [
        "Kategori Adı",
        "Ürün Sayısı",
        "Ortalama Fiyat (₺)",
        "Ort. Kalori (kcal)",
        "Ort. Protein (g)",
        "Ort. Karbonhidrat (g)",
        "Ort. Yağ (g)",
        "Ort. Şeker (g)"
    ]
    
    ws_summary.row_dimensions[3].height = 26
    for c_idx, s_h in enumerate(s_headers, 1):
        c = ws_summary.cell(row=3, column=c_idx, value=s_h)
        c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=gray_header, end_color=gray_header, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # Grupları hesapla
    categories = ["YİYECEK GRUBU", "TATLI GRUBU", "DONDURMA GRUBU", "ETİ GRUBU"]
    s_row = 4
    for cat in categories:
        cat_items = [x for x in DATA if x["kategori"] == cat]
        n_items = len(cat_items)
        avg_fiyat = sum(x["fiyat"] for x in cat_items) / n_items
        avg_kal = sum(x["kalori_kcal"] for x in cat_items) / n_items
        avg_prot = sum(x["protein_g"] for x in cat_items) / n_items
        avg_karb = sum(x["karbonhidrat_g"] for x in cat_items) / n_items
        avg_yag = sum(x["yag_g"] for x in cat_items) / n_items
        avg_seker = sum(x["seker_g"] for x in cat_items) / n_items
        
        ws_summary.row_dimensions[s_row].height = 22
        cat_fill = cat_colors.get(cat, "FFFFFF")
        
        vals = [cat, n_items, avg_fiyat, avg_kal, avg_prot, avg_karb, avg_yag, avg_seker]
        for c_idx, v in enumerate(vals, 1):
            c = ws_summary.cell(row=s_row, column=c_idx, value=v)
            c.font = Font(name="Segoe UI", size=9.5)
            c.border = thin_border
            c.fill = PatternFill(start_color=cat_fill, end_color=cat_fill, fill_type="solid")
            
            if c_idx == 1:
                c.font = Font(name="Segoe UI", size=9.5, bold=True)
                c.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx == 2:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '₺#,##0.00'
            elif c_idx == 4:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0'
                c.font = Font(name="Segoe UI", size=9.5, bold=True, color="C026D3")
            else:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0.0'
        s_row += 1

    # Özet Kolon Genişlikleri
    for col in ws_summary.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row == 1: continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len: max_len = len(val_str)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # -------------------------------------------------------------
    # 3. SAYFA: SAĞLIKLI BESLENME & DİYETİSYEN NOTLARI
    # -------------------------------------------------------------
    ws_notes = wb.create_sheet(title="Diyetisyen & Okul Rehberi")
    ws_notes.views.sheetView[0].showGridLines = True
    
    ws_notes.merge_cells("A1:F1")
    n_title = ws_notes["A1"]
    n_title.value = "FLORYA MEV KOLEJİ - SAĞLIKLI KANTİN TÜKETİM VE BESLENME REHBERİ"
    n_title.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    n_title.fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    n_title.alignment = Alignment(horizontal="center", vertical="center")
    ws_notes.row_dimensions[1].height = 32

    guide_points = [
        ("1. Yüksek Proteinli & Doyurucu Seçenekler:", "Izgara Köfte Sandviç (27.6g protein), Ayvalık Tost (25.4g protein) ve Ekmek Arası Şinitzel (24.2g protein) spor yapan ve gelişim çağındaki öğrenciler için en yüksek protein sağlayan seçeneklerdir."),
        ("2. Düşük Kalorili & Hafif Atıştırmalıklar:", "Taze Meyve (82 kcal), Twister Dondurma (68 kcal), Eti Puf (75 kcal) ve Eti Süt Burger (145 kcal) ara öğünlerde hafif ve dengeli enerji arayanlar için uygundur."),
        ("3. Lif Kaynağı ve Sindirim Dostu Ürünler:", "Eti Burçak Çörekotlu (5.2g lif), Eti Form (4.6g lif), Eti Lifalif Bar (3.4g lif) ve Çiğköfte Dürüm (6.2g lif) uzun süre tokluk hissi veren zengin lif kaynaklarıdır."),
        ("4. Şeker ve Enerji Dengesi (Önemli Not):", "Tatlı kupları (325-395 kcal) ve paketli çikolatalar yoğun şeker içerdiği için sınav veya yoğun antrenman öncesi haricinde haftada 1-2 kez ile sınırlandırılması önerilir."),
        ("5. İçecek Önerisi:", "Kantin alışverişlerinde bu ürünlerin yanında asitli ve yapay şekerli içecekler yerine mutlaka su, ayran veya sade süt tercih edilmelidir.")
    ]

    r_idx = 3
    for title, desc in guide_points:
        ws_notes.cell(row=r_idx, column=1, value=title).font = Font(name="Segoe UI", size=10.5, bold=True, color=navy_dark)
        ws_notes.cell(row=r_idx+1, column=1, value=desc).font = Font(name="Segoe UI", size=9.5, italic=True)
        ws_notes.row_dimensions[r_idx].height = 22
        ws_notes.row_dimensions[r_idx+1].height = 26
        r_idx += 3

    ws_notes.column_dimensions["A"].width = 110

    return wb

if __name__ == "__main__":
    wb = create_nutrition_workbook()
    
    file_name = "NOYA_Kantin_Urunleri_Besin_Degerleri_2026_2027.xlsx"
    
    # Hedef Konumlar
    destinations = [
        Path(r"c:\Users\doruk\Desktop\Florya_MEV_Operasyon") / file_name,
        Path(r"C:\Users\doruk\Desktop") / file_name,
        Path(r"G:\Drive'ım") / file_name,
        Path(r"G:\Drive'ım\NOYA") / file_name,
        Path(r"G:\Drive'ım\UYGULAMA WEB SİTESİ") / file_name,
        Path(r"G:\Drive'ım\Florya_MEV_Yedekler") / file_name,
        Path(r"c:\Users\doruk\Desktop\Florya_MEV_Operasyon\yedekler") / file_name
    ]
    
    # 1. Proje dizinine kaydet
    local_path = destinations[0]
    wb.save(local_path)
    print(f"[BASARILI] Yerel dosya olusturuldu: {local_path}")
    
    # 2. Diğer tüm hedeflere kopyala
    for dst in destinations[1:]:
        try:
            if dst.parent.exists():
                shutil.copy2(local_path, dst)
                print(f"[BASARILI] Kopyalandi: {dst}")
        except Exception as e:
            print(f"[UYARI] {dst} kopyalanamadi: {e}")
